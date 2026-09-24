#!/usr/bin/env python3
"""Capsem native artifact export adapter.

Rust owns the typed route, audit, and artifact lookup. This script is the
swappable VM tool adapter for document formats that should not be implemented
inside the Rust server.
"""

from __future__ import annotations

import json
import sys
from pathlib import Path
from typing import Any


ADAPTER = "python.export_artifact"


def main() -> int:
    if len(sys.argv) != 2:
        emit(
            ok=False,
            status="failed",
            artifact_id="",
            fmt="",
            message="usage: export_artifact.py <input.json>",
        )
        return 2

    request = json.loads(Path(sys.argv[1]).read_text())
    kind = require_string(request, "kind")
    fmt = require_string(request, "format")
    artifact = require_object(request, "artifact")
    artifacts = request.get("artifacts", [])
    output_dir = Path(require_string(request, "outputDir"))
    job_id = require_string(request, "jobId")
    output_dir.mkdir(parents=True, exist_ok=True)

    try:
        if kind == "spreadsheet" and fmt == "xlsx":
            return export_spreadsheet_xlsx(artifact, artifacts, output_dir, job_id)
        if kind == "slideDeck" and fmt == "html":
            return export_slide_deck_html(artifact, artifacts, output_dir, job_id)
        if kind == "slideDeck" and fmt in {"pptx", "pdf"}:
            return tool_unavailable(
                artifact,
                fmt,
                "slide deck export requires python-pptx, LibreOffice/OpenOffice, or another configured office toolchain",
            )
        return tool_unavailable(artifact, fmt, f"unsupported export kind/format: {kind}/{fmt}")
    except Exception as exc:
        emit(
            ok=False,
            status="failed",
            artifact_id=str(artifact.get("id", "")),
            fmt=fmt,
            message=str(exc),
        )
        return 1


def export_spreadsheet_xlsx(
    artifact: dict[str, Any],
    artifacts: list[Any],
    output_dir: Path,
    job_id: str,
) -> int:
    try:
        from openpyxl import Workbook
        from openpyxl.chart import BarChart, Reference
        from openpyxl.styles import Font
    except Exception as exc:
        return tool_unavailable(
            artifact,
            "xlsx",
            f"spreadsheet export requires openpyxl in the Capsem image: {exc}",
        )

    spec = require_object(artifact, "spec")
    columns = require_list(spec, "columns")
    rows = require_list(spec, "rows")

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = safe_sheet_title(str(artifact.get("title") or "Sheet"))

    for col_index, column in enumerate(columns, start=1):
        cell = sheet.cell(row=1, column=col_index, value=str(column))
        cell.font = Font(bold=True)

    for row_index, row in enumerate(rows, start=2):
        row_object = row if isinstance(row, dict) else {}
        for col_index, column in enumerate(columns, start=1):
            sheet.cell(row=row_index, column=col_index, value=row_object.get(str(column)))

    for width_index, column in enumerate(columns, start=1):
        sheet.column_dimensions[column_letter(width_index)].width = max(12, min(32, len(str(column)) + 4))

    add_first_matching_chart(workbook, sheet, artifact, artifacts, columns, len(rows))

    output_path = output_dir / f"{job_id}.xlsx"
    workbook.save(output_path)
    emit_file(artifact, "xlsx", output_path, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    return 0


def add_first_matching_chart(
    workbook: Any,
    sheet: Any,
    artifact: dict[str, Any],
    artifacts: list[Any],
    columns: list[Any],
    row_count: int,
) -> None:
    try:
        from openpyxl.chart import BarChart, Reference
    except Exception:
        return

    artifact_id = artifact.get("id")
    chart_artifact = next(
        (
            value
            for value in artifacts
            if isinstance(value, dict)
            and value.get("kind") == "chart"
            and isinstance(value.get("spec"), dict)
            and value["spec"].get("sourceArtifact") == artifact_id
        ),
        None,
    )
    if chart_artifact is None or row_count == 0:
        return

    chart_spec = chart_artifact["spec"]
    x_field = chart_spec.get("x")
    series = chart_spec.get("series") or []
    if not isinstance(series, list) or not series:
        return
    try:
        x_col = [str(value) for value in columns].index(str(x_field)) + 1
        y_col = [str(value) for value in columns].index(str(series[0]["field"])) + 1
    except Exception:
        return

    chart = BarChart()
    chart.title = str(chart_artifact.get("title") or "Chart")
    chart.y_axis.title = str(chart_spec.get("yLabel") or "")
    chart.x_axis.title = str(chart_spec.get("xLabel") or "")
    data = Reference(sheet, min_col=y_col, min_row=1, max_row=row_count + 1)
    cats = Reference(sheet, min_col=x_col, min_row=2, max_row=row_count + 1)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    chart_sheet = workbook.create_sheet("Charts")
    chart_sheet.add_chart(chart, "A1")


def export_slide_deck_html(
    artifact: dict[str, Any],
    artifacts: list[Any],
    output_dir: Path,
    job_id: str,
) -> int:
    artifact_by_id = {
        value.get("id"): value for value in artifacts if isinstance(value, dict) and value.get("id")
    }
    spec = require_object(artifact, "spec")
    slides = require_list(spec, "slides")
    sections: list[str] = []
    for slide_ref in slides:
        slide_artifact = artifact_by_id.get(slide_ref.get("artifactId")) if isinstance(slide_ref, dict) else None
        if not isinstance(slide_artifact, dict):
            continue
        slide_spec = slide_artifact.get("spec") if isinstance(slide_artifact.get("spec"), dict) else {}
        blocks = slide_spec.get("blocks") if isinstance(slide_spec.get("blocks"), list) else []
        body = "\n".join(render_slide_block(block, artifact_by_id) for block in blocks)
        sections.append(
            f"<section><h2>{escape(slide_artifact.get('title', 'Slide'))}</h2>{body}</section>"
        )
    html = f"""<!doctype html>
<html>
<head>
  <meta charset="utf-8">
  <title>{escape(artifact.get("title", "Slide Deck"))}</title>
  <style>
    body {{ font-family: system-ui, sans-serif; margin: 0; background: #f5f5f5; color: #111827; }}
    section {{ width: 960px; min-height: 540px; margin: 32px auto; padding: 48px; background: #fff; border: 1px solid #e5e7eb; }}
    h1, h2, h3 {{ margin-top: 0; }}
    table {{ border-collapse: collapse; width: 100%; }}
    th, td {{ border: 1px solid #d1d5db; padding: 8px; text-align: left; }}
  </style>
</head>
<body>
  <section><h1>{escape(artifact.get("title", "Slide Deck"))}</h1></section>
  {"".join(sections)}
</body>
</html>
"""
    output_path = output_dir / f"{job_id}.html"
    output_path.write_text(html)
    emit_file(artifact, "html", output_path, "text/html")
    return 0


def render_slide_block(block: Any, artifact_by_id: dict[str, dict[str, Any]]) -> str:
    if not isinstance(block, dict):
        return ""
    kind = block.get("kind")
    if kind == "text":
        return f"<article><h3>{escape(block.get('title', ''))}</h3><p>{escape(block.get('body', ''))}</p></article>"
    ref = artifact_by_id.get(block.get("artifactId"))
    if ref is None:
        return f"<p>Missing artifact: {escape(block.get('artifactId', ''))}</p>"
    spec = ref.get("spec") if isinstance(ref.get("spec"), dict) else {}
    if ref.get("kind") in {"table", "sheet"}:
        return render_table(spec)
    return f"<article><h3>{escape(ref.get('title', 'Artifact'))}</h3><pre>{escape(json.dumps(spec, indent=2))}</pre></article>"


def render_table(spec: dict[str, Any]) -> str:
    columns = spec.get("columns") if isinstance(spec.get("columns"), list) else []
    rows = spec.get("rows") if isinstance(spec.get("rows"), list) else []
    head = "".join(f"<th>{escape(column)}</th>" for column in columns)
    body_rows = []
    for row in rows:
        row_object = row if isinstance(row, dict) else {}
        body_rows.append("".join(f"<td>{escape(row_object.get(str(column), ''))}</td>" for column in columns))
    body = "".join(f"<tr>{row}</tr>" for row in body_rows)
    return f"<table><thead><tr>{head}</tr></thead><tbody>{body}</tbody></table>"


def tool_unavailable(artifact: dict[str, Any], fmt: str, message: str) -> int:
    emit(
        ok=False,
        status="toolUnavailable",
        artifact_id=str(artifact.get("id", "")),
        fmt=fmt,
        message=message,
    )
    return 0


def emit_file(artifact: dict[str, Any], fmt: str, path: Path, mime_type: str) -> None:
    emit(
        ok=True,
        status="complete",
        artifact_id=str(artifact.get("id", "")),
        fmt=fmt,
        file_path=str(path),
        mime_type=mime_type,
        bytes_value=path.stat().st_size,
    )


def emit(
    *,
    ok: bool,
    status: str,
    artifact_id: str,
    fmt: str,
    message: str | None = None,
    file_path: str | None = None,
    mime_type: str | None = None,
    bytes_value: int | None = None,
) -> None:
    print(
        json.dumps(
            {
                "ok": ok,
                "artifactId": artifact_id,
                "format": fmt,
                "adapter": ADAPTER,
                "status": status,
                "filePath": file_path,
                "mimeType": mime_type,
                "bytes": bytes_value,
                "message": message,
            }
        )
    )


def require_string(value: dict[str, Any], key: str) -> str:
    result = value.get(key)
    if not isinstance(result, str) or not result:
        raise ValueError(f"{key} must be a non-empty string")
    return result


def require_object(value: dict[str, Any], key: str) -> dict[str, Any]:
    result = value.get(key)
    if not isinstance(result, dict):
        raise ValueError(f"{key} must be an object")
    return result


def require_list(value: dict[str, Any], key: str) -> list[Any]:
    result = value.get(key)
    if not isinstance(result, list):
        raise ValueError(f"{key} must be a list")
    return result


def safe_sheet_title(value: str) -> str:
    cleaned = "".join("-" if ch in "[]:*?/\\\\" else ch for ch in value).strip()
    return (cleaned or "Sheet")[:31]


def column_letter(index: int) -> str:
    letters = ""
    while index:
        index, remainder = divmod(index - 1, 26)
        letters = chr(65 + remainder) + letters
    return letters


def escape(value: Any) -> str:
    return (
        str(value)
        .replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
        .replace('"', "&quot;")
    )


if __name__ == "__main__":
    raise SystemExit(main())
